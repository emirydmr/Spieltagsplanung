"""Liest Spielstätten aus einer Meldeliste und erstellt die Entfernungsmatrix.

Nutzung:
    python src/common/erstelle_distanzmatrix.py raw_data/Meldeliste_Jugend_Hinrunde_24-25.xlsx
"""

import csv
import sys
from pathlib import Path

import openpyxl

# Projekt-Root zum Python-Path hinzufügen
ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT))

from src.common.distanz import geocode_adressen, berechne_distanzmatrix


def extrahiere_spielstaetten(excel_path: str) -> dict[str, str]:
    """Extrahiert einzigartige Spielstätten aus der Meldeliste.

    Returns:
        Dict: voller Spielstätten-String -> Adresse (ohne Facility-Name)
    """
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb[wb.sheetnames[0]]

    spielstaetten: dict[str, str] = {}

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3:
            continue
        # Spielfeld-Spalten: 21, 23, 25, 27 (0-indexed)
        for col_idx in [21, 23, 25, 27]:
            if len(row) > col_idx:
                val = row[col_idx]
                if val and str(val).strip():
                    full = str(val).strip()
                    parts = full.split(", ")
                    if len(parts) >= 2:
                        adresse = ", ".join(parts[1:])
                        spielstaetten[full] = adresse

    wb.close()
    return spielstaetten


def main():
    if len(sys.argv) < 2:
        print("Nutzung: python erstelle_distanzmatrix.py <meldeliste.xlsx>")
        sys.exit(1)

    excel_path = sys.argv[1]
    print(f"Lese Spielstätten aus: {excel_path}")

    # 1. Spielstätten extrahieren
    spielstaetten = extrahiere_spielstaetten(excel_path)
    adressen = sorted(set(spielstaetten.values()))
    print(f"  {len(spielstaetten)} Spielstätten, {len(adressen)} einzigartige Adressen\n")

    # 2. Geocoding
    print("Starte Geocoding (Nominatim/OpenStreetMap)...")
    print("  (Gecachte Adressen werden übersprungen)\n")
    koordinaten = geocode_adressen(adressen)

    gefunden = len(koordinaten)
    nicht_gefunden = len(adressen) - gefunden
    print(f"\nErgebnis: {gefunden}/{len(adressen)} Adressen geocodiert")
    if nicht_gefunden > 0:
        fehlend = set(adressen) - set(koordinaten.keys())
        print(f"  Nicht gefunden ({nicht_gefunden}):")
        for a in sorted(fehlend):
            print(f"    - {a}")

    # 3. Distanzmatrix berechnen
    print("\nBerechne Distanzmatrix...")
    matrix = berechne_distanzmatrix(koordinaten)

    # 4. Als CSV speichern
    output_path = ROOT / "config" / "distanzmatrix.csv"
    keys = sorted(matrix.keys())
    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow([""] + keys)
        for a in keys:
            row = [a] + [str(matrix[a].get(b, "")) for b in keys]
            writer.writerow(row)

    print(f"Distanzmatrix gespeichert: {output_path}")
    print(f"  {len(keys)} x {len(keys)} = {len(keys)**2} Einträge")

    # 5. Kurzstatistik
    alle_distanzen = [matrix[a][b] for a in keys for b in keys if a != b]
    if alle_distanzen:
        print(f"\n  Min: {min(alle_distanzen):.1f} km")
        print(f"  Max: {max(alle_distanzen):.1f} km")
        print(f"  Durchschnitt: {sum(alle_distanzen)/len(alle_distanzen):.1f} km")


if __name__ == "__main__":
    main()
