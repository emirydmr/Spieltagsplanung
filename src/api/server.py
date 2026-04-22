"""FastAPI Backend für die Spieltagsplanung.

Starten:
    python -m uvicorn src.api.server:app --reload --port 8000
"""

import sys
import os

# Force UTF-8 stdout/stderr on Windows (prevents charmap codec errors)
if sys.platform == "win32":
    os.environ.setdefault("PYTHONIOENCODING", "utf-8")
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import tempfile
import shutil
import json
import time as _time
from datetime import datetime
from pathlib import Path

from fastapi import FastAPI, UploadFile, File, HTTPException, Request
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT))

from src.data_import.meldeliste_parser import parse_meldeliste, verknuepfe_koordinaten
from src.data_import.rueckrunde_parser import (
    parse_rueckrunde_einteilung, match_teams_gegen_meldeliste, staffeln_to_einteilung_result,
)
from src.staffeleinteilung.algorithmus import gruppiere_mannschaften, einteilung_erstellen
from src.staffeleinteilung.scoring import ScoreGewichte
from src.common.distanz import haversine_km
from src.spielplanerstellung.wuensche_parser import parse_wuensche_llm
from src.spielplanerstellung.wuensche import Wunsch, WunschKategorie, WunschPrio, VereinsWuensche
from src.spielplanerstellung.spielplan import generiere_alle_spielplaene, spielplan_to_dict

app = FastAPI(title="Spieltagsplaner", version="1.0")

# Verzeichnis für gespeicherte Spielpläne
SPIELPLAN_DIR = ROOT / "spielplan_logs"
SPIELPLAN_DIR.mkdir(exist_ok=True)

# Static files
UI_DIR = ROOT / "ui"
app.mount("/assets", StaticFiles(directory=str(UI_DIR / "assets")), name="assets")


# ─── Pages ─────────────────────────────────────────────────────

@app.get("/")
async def index():
    return FileResponse(str(UI_DIR / "index.html"))


@app.get("/app")
async def app_page():
    return FileResponse(str(UI_DIR / "app.html"))


@app.get("/builder")
async def builder_page():
    return FileResponse(str(UI_DIR / "builder.html"))


# ─── API ───────────────────────────────────────────────────────

@app.post("/api/einteilung")
async def api_einteilung(
    file: UploadFile = File(...),
    w_distanz: float = 1.0,
    w_region: float = 50.0,
    w_balance: float = 20.0,
    max_staffel_size: int = 11,
):
    """Upload Meldeliste Excel -> Staffeleinteilung berechnen."""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Nur Excel-Dateien (.xlsx) erlaubt")

    # Save uploaded file to temp
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        shutil.copyfileobj(file.file, tmp)
        tmp_path = tmp.name

    try:
        mannschaften = parse_meldeliste(tmp_path)
        n_coords = verknuepfe_koordinaten(mannschaften)

        gewichte = ScoreGewichte(
            w_distanz=w_distanz,
            w_region=w_region,
            w_balance=w_balance,
            max_staffel_size=max_staffel_size,
        )

        gruppen = gruppiere_mannschaften(mannschaften)

        result = {
            "total_teams": len(mannschaften),
            "teams_mit_coords": n_coords,
            "gruppen": [],
        }

        for (ak, sk, topf), ms in sorted(gruppen.items()):
            topf_str = f"Topf {topf}" if topf > 0 else "fix"
            print(f"[Einteilung] {ak} {sk} {topf_str} – {len(ms)} Teams ...", flush=True)
            t0 = _time.time()

            gruppe_data = {
                "altersklasse": ak,
                "spielklasse": sk,
                "topf": topf_str,
                "n_teams": len(ms),
                "staffeln": [],
                "score": None,
                "merged": False,
            }

            # Regionenstaffel: single staffel if small enough
            if topf == 0 and len(ms) <= max_staffel_size:
                gruppe_data["staffeln"] = [_staffel_to_dict(ms)]
                result["gruppen"].append(gruppe_data)
                continue

            if topf == 0 and len(ms) > max_staffel_size:
                gruppe_data["merged"] = True

            if len(ms) < 4:
                staffeln_list = [ms]
            else:
                staffeln_list, score = einteilung_erstellen(ms, gewichte)
                gruppe_data["score"] = {
                    "total": round(score.total, 1),
                    "distanz": round(score.distanz_score, 1),
                    "region": int(score.region_score),
                    "balance": round(score.balance_score, 2),
                    "violations": score.hard_violations,
                }

            for staffel in staffeln_list:
                gruppe_data["staffeln"].append(_staffel_to_dict(staffel))

            result["gruppen"].append(gruppe_data)
            print(f"[Einteilung] {ak} {sk} {topf_str} – fertig in {_time.time() - t0:.1f}s", flush=True)

        return JSONResponse(result)

    except Exception as e:
        raise HTTPException(500, f"Fehler bei Verarbeitung: {str(e)}")
    finally:
        Path(tmp_path).unlink(missing_ok=True)


def _staffel_to_dict(mannschaften: list) -> dict:
    """Konvertiert eine Staffel (Liste von Mannschaft) in ein JSON-fähiges dict."""
    teams = []
    for m in sorted(mannschaften, key=lambda x: x.mannschaftsname):
        team = {
            "mannschaft": m.mannschaftsname,
            "verein": m.vereinsname,
            "verein_nr": m.verein_nr,
            "region": m.bezirk_alt or "?",
            "topf": m.topf,
            "wuensche_text": m.wuensche_text or "",
        }
        if m.spielstaette:
            team["spielstaette"] = m.spielstaette.name or ""
            team["adresse"] = m.spielstaette.adresse or ""
            team["lat"] = m.spielstaette.lat
            team["lon"] = m.spielstaette.lon
        teams.append(team)

    # Max pairwise distance
    max_dist = 0
    for i in range(len(mannschaften)):
        a = mannschaften[i]
        if not a.spielstaette or a.spielstaette.lat is None:
            continue
        for j in range(i + 1, len(mannschaften)):
            b = mannschaften[j]
            if not b.spielstaette or b.spielstaette.lat is None:
                continue
            d = haversine_km(
                a.spielstaette.lat, a.spielstaette.lon,
                b.spielstaette.lat, b.spielstaette.lon,
            )
            if d > max_dist:
                max_dist = d

    doppelrunde = len(mannschaften) < 5

    return {
        "n_teams": len(mannschaften),
        "teams": teams,
        "max_distanz_km": round(max_dist, 1),
        "doppelrunde": doppelrunde,
    }


# ─── Rückrunde Einteilung ─────────────────────────────────────

@app.post("/api/einteilung/rueckrunde")
async def api_einteilung_rueckrunde(
    meldeliste: UploadFile = File(...),
    einteilung: UploadFile = File(...),
):
    """Rückrunde: Liest vorgegebene Staffelzuordnungen aus der Einteilungs-Excel.

    Erwartet zwei Dateien:
      - meldeliste: DFBnet-Meldeliste (Spielstätten, Wünsche)
      - einteilung: Einteilungs-Excel mit Staffelzuordnungen vom Spielleiter
    """
    for f in (meldeliste, einteilung):
        if not f.filename.endswith((".xlsx", ".xls")):
            raise HTTPException(400, f"Nur Excel-Dateien erlaubt: {f.filename}")

    tmp_ml = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp_et = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    try:
        shutil.copyfileobj(meldeliste.file, tmp_ml)
        tmp_ml.close()
        shutil.copyfileobj(einteilung.file, tmp_et)
        tmp_et.close()

        # 1. Meldeliste parsen (Spielstätten, Wünsche, Koordinaten)
        mannschaften = parse_meldeliste(tmp_ml.name)
        n_coords = verknuepfe_koordinaten(mannschaften)

        # 2. Rückrunde-Einteilung parsen
        staffeln = parse_rueckrunde_einteilung(tmp_et.name)

        # 3. Teams gegen Meldeliste matchen
        matched, unmatched = match_teams_gegen_meldeliste(staffeln, mannschaften)
        print(f"[Rückrunde] {matched} Teams gematcht, {unmatched} ohne Match")

        # 4. Meldeliste-Teams ohne Match finden (Nachmeldungen / neue Teams)
        matched_ids = set()
        for s in staffeln:
            for t in s.teams:
                if t.mannschaft:
                    matched_ids.add(id(t.mannschaft))
        neue_teams = []
        for m in mannschaften:
            if id(m) not in matched_ids:
                t = {"mannschaft": m.mannschaftsname, "verein": m.vereinsname,
                     "verein_nr": m.verein_nr, "altersklasse": m.altersklasse,
                     "region": m.bezirk_alt or "?"}
                if m.spielstaette:
                    t["lat"] = m.spielstaette.lat
                    t["lon"] = m.spielstaette.lon
                neue_teams.append(t)

        # 5. In API-Format konvertieren
        result = staffeln_to_einteilung_result(staffeln, len(mannschaften), n_coords)
        result["matched"] = matched
        result["unmatched"] = unmatched
        result["neue_teams"] = neue_teams

        return JSONResponse(result)

    except Exception as e:
        raise HTTPException(500, f"Fehler bei Rückrunde-Verarbeitung: {str(e)}")
    finally:
        Path(tmp_ml.name).unlink(missing_ok=True)
        Path(tmp_et.name).unlink(missing_ok=True)


# ─── Builder: Team-Pool ───────────────────────────────────────

@app.post("/api/builder/teams")
async def api_builder_teams(meldeliste: UploadFile = File(...)):
    """Parst die Meldeliste und gibt einen flachen Team-Pool zurück.

    Jedes Team bekommt eine eindeutige _id und alle Infos die der
    Staffel-Builder braucht (Name, Region, Koordinaten, Hinrunde-Daten).
    """
    if not meldeliste.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Nur Excel-Dateien (.xlsx) erlaubt")

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        shutil.copyfileobj(meldeliste.file, tmp)
        tmp_path = tmp.name

    try:
        mannschaften = parse_meldeliste(tmp_path)
        n_coords = verknuepfe_koordinaten(mannschaften)

        teams = []
        for i, m in enumerate(mannschaften):
            team = {
                "_id": f"t_{i}",
                "mannschaft": m.mannschaftsname,
                "verein": m.vereinsname,
                "verein_nr": m.verein_nr,
                "altersklasse": m.altersklasse,
                "region": m.bezirk_alt or "?",
                "topf": m.topf,
                "wuensche_text": m.wuensche_text or "",
            }
            if m.spielstaette:
                team["spielstaette"] = m.spielstaette.name or ""
                team["adresse"] = m.spielstaette.adresse or ""
                team["lat"] = m.spielstaette.lat
                team["lon"] = m.spielstaette.lon
            # Hinrunde-Daten (leer bei Erstimport)
            team["rang_hinrunde"] = None
            team["punkte_hinrunde"] = None
            team["quotient_hinrunde"] = None
            teams.append(team)

        return JSONResponse({
            "teams": teams,
            "total": len(teams),
            "teams_mit_coords": n_coords,
        })

    except Exception as e:
        raise HTTPException(500, f"Fehler beim Parsen: {str(e)}")
    finally:
        Path(tmp_path).unlink(missing_ok=True)


# ─── Excel Export ──────────────────────────────────────────────

@app.post("/api/export")
async def api_export(request: Request):
    """Generiert eine Excel-Datei aus dem Einteilungsergebnis."""
    data = await request.json()
    gruppen = data.get("gruppen", [])
    if not gruppen:
        raise HTTPException(400, "Keine Gruppen vorhanden")

    wb = Workbook()

    # ── Styles ──
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="C41230", end_color="C41230", fill_type="solid")
    staffel_font = Font(bold=True, size=11, color="C41230")
    thin_border = Border(
        bottom=Side(style="thin", color="E5E7EB"),
    )
    center = Alignment(horizontal="center")

    # ── Sheet 1: Übersicht ──
    ws_ueb = wb.active
    ws_ueb.title = "Übersicht"
    ueb_headers = ["Altersklasse", "Topf", "Teams", "Staffeln", "Ø Distanz (km)", "Probleme"]
    for ci, h in enumerate(ueb_headers, 1):
        cell = ws_ueb.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for ri, g in enumerate(gruppen, 2):
        ws_ueb.cell(row=ri, column=1, value=g["altersklasse"])
        ws_ueb.cell(row=ri, column=2, value=g["topf"])
        ws_ueb.cell(row=ri, column=3, value=g["n_teams"]).alignment = center
        ws_ueb.cell(row=ri, column=4, value=len(g["staffeln"])).alignment = center
        dist = g["score"]["distanz"] if g.get("score") else "-"
        ws_ueb.cell(row=ri, column=5, value=dist).alignment = center
        viol = g["score"]["violations"] if g.get("score") else 0
        ws_ueb.cell(row=ri, column=6, value=viol).alignment = center

    for col in ws_ueb.columns:
        ws_ueb.column_dimensions[col[0].column_letter].width = 18

    # ── Sheets per Altersklasse/Topf ──
    for g in gruppen:
        sheet_name = f"{g['altersklasse']} {g['topf']}"[:31]  # Excel max 31 chars
        ws = wb.create_sheet(title=sheet_name)

        row = 1
        for si, staffel in enumerate(g["staffeln"]):
            # Staffel header
            label = f"Staffel {si + 1}"
            if staffel.get("doppelrunde"):
                label += " (Doppelrunde)"
            label += f"  —  {staffel['n_teams']} Teams, Max. {staffel['max_distanz_km']} km"
            cell = ws.cell(row=row, column=1, value=label)
            cell.font = staffel_font
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            row += 1

            # Table headers
            team_headers = ["Nr.", "Mannschaft", "Verein", "Region", "Ort"]
            for ci, h in enumerate(team_headers, 1):
                cell = ws.cell(row=row, column=ci, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
            row += 1

            # Team rows
            for ti, t in enumerate(staffel["teams"], 1):
                ws.cell(row=row, column=1, value=ti).alignment = center
                ws.cell(row=row, column=2, value=t["mannschaft"])
                ws.cell(row=row, column=3, value=t.get("verein", ""))
                ws.cell(row=row, column=4, value=t.get("region", ""))
                ort = t.get("adresse", "").split(", ")[-1] if t.get("adresse") else ""
                ws.cell(row=row, column=5, value=ort)
                for ci in range(1, 6):
                    ws.cell(row=row, column=ci).border = thin_border
                row += 1

            row += 1  # empty row between staffeln

        # Column widths
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 14
        ws.column_dimensions["E"].width = 22

    # Save to temp file
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name)
    tmp.close()

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="Staffeleinteilung.xlsx",
    )


# ─── Rückrunde Excel Export (Jochen-Format) ────────────────────

@app.post("/api/export/rueckrunde")
async def api_export_rueckrunde(request: Request):
    """Exportiert die Einteilung im Jochen-Format (gleiche Struktur wie Eingangs-Excel).

    Layout pro Sheet (= Altersklasse):
    - Links (A-O): Staffel-Blöcke nebeneinander (max 3 pro Zeile)
      je Block: Rang | Mannschaft | Punkte | Q | (leer)
    - Rechts (P-Y): Gesamtübersicht nach Region (Unterland / Hohenlohe)
      mit LS/KS/BS-Markierung
    """
    data = await request.json()
    gruppen = data.get("gruppen", [])
    saison = data.get("saison", "2025/26")
    if not gruppen:
        raise HTTPException(400, "Keine Gruppen vorhanden")

    wb = _build_rueckrunde_workbook(gruppen, saison)

    # Einteilung als JSON sichern (für Spielplan-Generierung Rückrunde)
    rr_einteilung = {"gruppen": gruppen, "saison": saison, "total_teams": sum(g.get("n_teams", 0) for g in gruppen)}
    rr_path = ROOT / "config" / "rueckrunde_einteilung.json"
    with open(rr_path, "w", encoding="utf-8") as f:
        json.dump(rr_einteilung, f, ensure_ascii=False, indent=1)
    print(f"[RR] Einteilung gespeichert: {rr_path.name} ({rr_einteilung['total_teams']} Teams)")

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name)
    tmp.close()

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"Einteilung_Rueckrunde_{saison.replace('/', '_')}.xlsx",
    )


@app.get("/api/rueckrunde/einteilung")
async def api_get_rueckrunde_einteilung():
    """Gibt die gespeicherte Rückrunde-Einteilung zurück (falls vorhanden)."""
    rr_path = ROOT / "config" / "rueckrunde_einteilung.json"
    if not rr_path.exists():
        raise HTTPException(404, "Keine Rückrunde-Einteilung vorhanden")
    with open(rr_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    return JSONResponse(data)


def _build_rueckrunde_workbook(gruppen: list[dict], saison: str) -> Workbook:
    """Baut eine Excel-Datei im Jochen-Format."""
    from datetime import date as _date
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # ── Styles (Jochen-Format: Arial Narrow) ──
    title_font = Font(name="Arial Narrow", bold=True, size=22, underline="single")
    header_font = Font(name="Arial Narrow", bold=True, size=16)
    subheader_font = Font(name="Arial Narrow", bold=True, size=12)
    team_font = Font(name="Arial Narrow", size=14)
    team_font_small = Font(name="Arial Narrow", size=11)
    marker_font = Font(name="Arial Narrow", bold=True, size=11)
    stand_font = Font(name="Arial Narrow", size=11, color="888888")

    center = Alignment(horizontal="center", vertical="center")

    # Farben für Staffeltypen
    ls_fill = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
    ks_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    # Gruppiere nach Altersklasse
    ak_gruppen: dict[str, list[dict]] = {}
    for g in gruppen:
        ak = g["altersklasse"]
        ak_gruppen.setdefault(ak, []).append(g)

    first_sheet = True
    for ak in sorted(ak_gruppen.keys()):
        if first_sheet:
            ws = wb.active
            ws.title = ak[:31]
            first_sheet = False
        else:
            ws = wb.create_sheet(title=ak[:31])

        # Spaltenbreiten (Jochen-Format)
        widths = [4.6, 50, 7.9, 5, 6.4] * 3
        widths += [8.6, 4.4, 43, 7, 7, 11.4, 4.4, 43, 7, 7, 11.4]
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w

        # ── Titel-Zeile ──
        ws.cell(row=1, column=1, value=f"Vorschlag Einteilung Rückrunde {ak} {saison}").font = title_font
        ws.merge_cells("A1:G1")
        today = _date.today()
        ws.cell(row=1, column=12, value=f"Stand: {today.strftime('%d.%m.%y')}").font = stand_font

        # Sammle alle Staffeln dieser AK, gruppiert nach Typ
        ak_staffeln = []
        for g in ak_gruppen[ak]:
            for s in g.get("staffeln", []):
                ak_staffeln.append({"staffel": s, "gruppe": g})

        # Sortiere: Bezirksstaffel -> Leistungsstaffel -> Kreisstaffel
        type_order = {"Bezirksstaffel": 0, "Leistungsstaffel": 1, "Kreisstaffel": 2}
        ak_staffeln.sort(key=lambda x: (
            type_order.get(x["gruppe"]["topf"], 3),
            x["staffel"].get("staffel_name", ""),
        ))

        # ── Links: Staffel-Blöcke (max 3 nebeneinander) ──
        row = 4
        staffel_idx = 0
        current_type = None

        while staffel_idx < len(ak_staffeln):
            entry = ak_staffeln[staffel_idx]
            typ = entry["gruppe"]["topf"]

            if typ != current_type:
                if current_type is not None:
                    row += 2
                current_type = typ

            # Bis zu 3 Staffeln nebeneinander
            block_start_idx = staffel_idx
            cols_used = 0
            while staffel_idx < len(ak_staffeln) and cols_used < 3:
                if ak_staffeln[staffel_idx]["gruppe"]["topf"] != current_type:
                    break
                staffel_idx += 1
                cols_used += 1

            block_staffeln = ak_staffeln[block_start_idx:staffel_idx]

            # Header-Zeile
            for bi, entry in enumerate(block_staffeln):
                col_offset = bi * 5 + 1
                s = entry["staffel"]
                name = s.get("staffel_name", f"Staffel {bi + 1}")
                dr = " (Doppelrunde)" if s.get("doppelrunde") else ""

                ws.cell(row=row, column=col_offset, value=f"{name}{dr}").font = header_font
                ws.merge_cells(
                    start_row=row, start_column=col_offset,
                    end_row=row, end_column=col_offset + 1,
                )
                ws.cell(row=row, column=col_offset + 2, value="Punkte").font = subheader_font
                ws.cell(row=row, column=col_offset + 2).alignment = center
                ws.cell(row=row, column=col_offset + 3, value="Q").font = subheader_font
                ws.cell(row=row, column=col_offset + 3).alignment = center

            row += 1

            # Teams
            max_teams = max(len(e["staffel"].get("teams", [])) for e in block_staffeln)
            for ti in range(max_teams):
                for bi, entry in enumerate(block_staffeln):
                    col_offset = bi * 5 + 1
                    teams = entry["staffel"].get("teams", [])
                    if ti >= len(teams):
                        continue
                    t = teams[ti]
                    rang = t.get("rang_hinrunde", ti + 1)
                    ws.cell(row=row, column=col_offset, value=f"{rang}.").font = team_font
                    ws.cell(row=row, column=col_offset).alignment = center
                    ws.cell(row=row, column=col_offset + 1, value=t.get("mannschaft", "")).font = team_font
                    punkte = t.get("punkte_hinrunde")
                    if punkte is not None:
                        ws.cell(row=row, column=col_offset + 2, value=punkte).font = team_font
                        ws.cell(row=row, column=col_offset + 2).alignment = center
                    quotient = t.get("quotient_hinrunde")
                    if quotient is not None:
                        q_val = round(quotient, 4) if isinstance(quotient, float) else quotient
                        ws.cell(row=row, column=col_offset + 3, value=q_val).font = team_font
                        ws.cell(row=row, column=col_offset + 3).alignment = center
                row += 1

        # ── Rechts: Regionale Zusammenfassung (P-Y) ──
        _write_regional_summary(ws, ak_staffeln, team_font_small, subheader_font, marker_font, center)

    return wb


def _write_regional_summary(ws, ak_staffeln, team_font, subheader_font, marker_font, center):
    """Schreibt die regionale Zusammenfassung rechts (Spalten P-Y)."""
    regions: dict[str, list[dict]] = {}

    for entry in ak_staffeln:
        typ = entry["gruppe"]["topf"]
        marker = "LS" if "Leistung" in typ else "BS" if "Bezirk" in typ else "KS"

        for t in entry["staffel"].get("teams", []):
            region = t.get("region", "?")
            team_data = {
                "mannschaft": t.get("mannschaft", ""),
                "punkte": t.get("punkte_hinrunde"),
                "quotient": t.get("quotient_hinrunde"),
                "marker": marker,
            }
            if "Unterland" in region:
                regions.setdefault("Unterland", []).append(team_data)
            elif "Hohenlohe" in region:
                regions.setdefault("Hohenlohe", []).append(team_data)
            else:
                regions.setdefault("Sonstige", []).append(team_data)

    for key in regions:
        regions[key].sort(key=lambda x: (-(x["quotient"] or 0), -(x["punkte"] or 0)))

    # Unterland ab Spalte Q (17), Hohenlohe ab V (22)
    col_ul, col_hl = 17, 22

    for label, col in [("Unterland", col_ul), ("Hohenlohe", col_hl)]:
        ws.cell(row=4, column=col, value=label).font = subheader_font
        ws.cell(row=4, column=col + 1, value="Punkte").font = subheader_font
        ws.cell(row=4, column=col + 1).alignment = center
        ws.cell(row=4, column=col + 2, value="Q").font = subheader_font
        ws.cell(row=4, column=col + 2).alignment = center

        for i, t in enumerate(regions.get(label, [])):
            r = 5 + i
            ws.cell(row=r, column=col - 1, value=f"{i + 1}.").font = team_font
            ws.cell(row=r, column=col - 1).alignment = center
            ws.cell(row=r, column=col, value=t["mannschaft"]).font = team_font
            if t["punkte"] is not None:
                ws.cell(row=r, column=col + 1, value=t["punkte"]).font = team_font
                ws.cell(row=r, column=col + 1).alignment = center
            if t["quotient"] is not None:
                q_val = round(t["quotient"], 4) if isinstance(t["quotient"], float) else t["quotient"]
                ws.cell(row=r, column=col + 2, value=q_val).font = team_font
                ws.cell(row=r, column=col + 2).alignment = center
            ws.cell(row=r, column=col + 3, value=t["marker"]).font = marker_font
            ws.cell(row=r, column=col + 3).alignment = center


# ─── Vereinswünsche LLM Parsing ───────────────────────────────

@app.post("/api/wuensche/parse")
async def api_parse_wuensche(request: Request):
    """Freitext-Wünsche per LLM in strukturierte Flags umwandeln."""
    data = await request.json()
    freitext = data.get("text", "")
    mannschaft = data.get("mannschaft", "")
    verein = data.get("verein", "")
    provider = data.get("provider", "ollama")
    model = data.get("model", "llama3")
    api_base = data.get("api_base")
    api_key = data.get("api_key")

    if not freitext.strip():
        raise HTTPException(400, "Kein Text angegeben")

    try:
        result = parse_wuensche_llm(
            freitext=freitext,
            mannschaft=mannschaft,
            verein=verein,
            provider=provider,
            model=model,
            api_base=api_base,
            api_key=api_key,
        )
        return JSONResponse({
            "mannschaft": result.mannschaft,
            "verein": result.verein,
            "wuensche": [
                {
                    "kategorie": w.kategorie.value,
                    "prioritaet": w.prioritaet.value,
                    "beschreibung": w.beschreibung,
                    "datum": w.datum,
                    "wochentag": w.wochentag,
                    "uhrzeit": w.uhrzeit,
                    "bezug_mannschaft": w.bezug_mannschaft,
                }
                for w in result.wuensche
            ],
        })
    except Exception as e:
        raise HTTPException(500, f"LLM-Fehler: {str(e)}")


# ─── Spielplan-Generierung ─────────────────────────────────────

@app.post("/api/spielplan")
async def api_spielplan(request: Request):
    """Generiert Spielpläne für alle Staffeln aus der Einteilung."""
    data = await request.json()
    einteilung = data.get("einteilung")
    saison = data.get("saison", "")
    sperrtage_raw = data.get("sperrtage", [])
    if not einteilung or not einteilung.get("gruppen"):
        raise HTTPException(400, "Keine Einteilung vorhanden")

    # Sperrtage parsen (ISO-Strings -> date-Objekte)
    from datetime import date as _date
    sperrtage = set()
    for s in (sperrtage_raw or []):
        try:
            sperrtage.add(_date.fromisoformat(s))
        except (ValueError, TypeError):
            pass

    try:
        # Wünsche regelbasiert parsen (schnell, kein LLM)
        wuensche = _parse_wuensche_fast(einteilung, saison)

        plaene = generiere_alle_spielplaene(
            einteilung, wuensche=wuensche if wuensche else None,
            sperrtage=sperrtage if sperrtage else None,
        )
        result = {
            "spielplaene": [spielplan_to_dict(p) for p in plaene],
            "total_staffeln": len(plaene),
            "wuensche_parsed": len(wuensche),
        }

        # Spielplan persistent speichern
        _save_spielplan_log(result, saison)

        return JSONResponse(result)
    except Exception as e:
        raise HTTPException(500, f"Fehler bei Spielplan-Generierung: {str(e)}")


def _save_spielplan_log(result: dict, saison: str) -> Path:
    """Speichert das Spielplan-Ergebnis als JSON-Datei mit Metadaten."""
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    total_konflikte = sum(
        p.get("score", {}).get("platz_konflikte", 0)
        for p in result.get("spielplaene", []) if p.get("score")
    )
    total_wunsch = sum(
        p.get("score", {}).get("wunsch_verletzungen", 0)
        for p in result.get("spielplaene", []) if p.get("score")
    )
    total_spiele = sum(
        len(s["spiele"])
        for p in result.get("spielplaene", []) for s in p["spieltage"]
    )
    log_entry = {
        "timestamp": datetime.now().isoformat(),
        "saison": saison,
        "total_staffeln": result.get("total_staffeln", 0),
        "total_spiele": total_spiele,
        "total_konflikte": total_konflikte,
        "total_wunsch_verletzungen": total_wunsch,
        "wuensche_parsed": result.get("wuensche_parsed", 0),
        "spielplaene": result.get("spielplaene", []),
    }

    filename = f"spielplan_{ts}.json"
    filepath = SPIELPLAN_DIR / filename
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(log_entry, f, ensure_ascii=False, indent=1)
    print(f"[LOG] Spielplan gespeichert: {filepath.name} "
          f"({total_spiele} Spiele, {total_konflikte} Platz-Konflikte, {total_wunsch} Wunsch-Verl.)")
    return filepath


@app.get("/api/spielplan/history")
async def api_spielplan_history():
    """Gibt eine Liste aller gespeicherten Spielpläne zurück (ohne Detaildaten)."""
    entries = []
    for fp in sorted(SPIELPLAN_DIR.glob("spielplan_*.json"), reverse=True):
        try:
            with open(fp, "r", encoding="utf-8") as f:
                data = json.load(f)
            entries.append({
                "filename": fp.name,
                "timestamp": data.get("timestamp", ""),
                "saison": data.get("saison", ""),
                "total_staffeln": data.get("total_staffeln", 0),
                "total_spiele": data.get("total_spiele", 0),
                "total_konflikte": data.get("total_konflikte", 0),
                "wuensche_parsed": data.get("wuensche_parsed", 0),
            })
        except (json.JSONDecodeError, KeyError):
            continue
    return JSONResponse({"history": entries})


@app.get("/api/spielplan/history/{filename}")
async def api_spielplan_load(filename: str):
    """Lädt einen bestimmten gespeicherten Spielplan."""
    filepath = SPIELPLAN_DIR / filename
    if not filepath.exists() or not filepath.name.startswith("spielplan_"):
        raise HTTPException(404, "Spielplan nicht gefunden")
    with open(filepath, "r", encoding="utf-8") as f:
        data = json.load(f)
    return JSONResponse(data)


# ─── Manuelle Spielplan-Bearbeitung ────────────────────────────

@app.post("/api/spielplan/edit")
async def api_spielplan_edit(request: Request):
    """Wendet manuelle Änderungen auf Spielpläne an und berechnet Konflikte neu.

    Erwartet JSON mit:
      - spielplaene: aktueller Stand aller Spielpläne
      - edits: Liste von Änderungen [{staffel_idx, spieltag_nr, spiel_idx, field, value}]

    field kann sein: 'datum', 'anstosszeit', 'swap_teams'
    """
    data = await request.json()
    spielplaene = data.get("spielplaene", [])
    edits = data.get("edits", [])

    if not spielplaene:
        raise HTTPException(400, "Keine Spielpläne vorhanden")
    if not edits:
        raise HTTPException(400, "Keine Änderungen angegeben")

    for edit in edits:
        si = edit.get("staffel_idx")
        st_nr = edit.get("spieltag_nr")
        sp_idx = edit.get("spiel_idx")
        fld = edit.get("field", "")
        val = edit.get("value", "")

        # Staffel finden
        plan = None
        for p in spielplaene:
            if p["staffel_idx"] == si:
                plan = p
                break
        if plan is None:
            continue

        # Spieltag finden
        spieltag = None
        for st in plan["spieltage"]:
            if st["nummer"] == st_nr:
                spieltag = st
                break
        if spieltag is None:
            continue

        # Spiel finden
        if sp_idx < 0 or sp_idx >= len(spieltag["spiele"]):
            continue
        spiel = spieltag["spiele"][sp_idx]

        # Änderung anwenden
        if fld == "datum":
            spiel["datum"] = val
        elif fld == "anstosszeit":
            spiel["anstosszeit"] = val
        elif fld == "swap_teams":
            spiel["heim"], spiel["gast"] = spiel["gast"], spiel["heim"]
            # Spielfeld wechselt zum neuen Heim-Team
            if val:
                spiel["spielfeld"] = val

    # Konflikte neu berechnen
    _recalculate_konflikte(spielplaene)

    return JSONResponse({"spielplaene": spielplaene})


def _recalculate_konflikte(spielplaene: list[dict]) -> None:
    """Berechnet platz_konflikte für JSON-Spielpläne neu."""
    from collections import defaultdict

    _AK_HALBFELD = {"F-Junioren", "F-Juniorinnen", "E-Junioren", "E-Juniorinnen",
                     "D-Junioren", "D-Juniorinnen", "Bambini"}
    _AK_DAUER = {
        "Bambini": 40, "F-Junioren": 40, "F-Juniorinnen": 40,
        "E-Junioren": 50, "E-Juniorinnen": 50,
        "D-Junioren": 60, "D-Juniorinnen": 60,
        "C-Junioren": 70, "C-Juniorinnen": 70,
        "B-Junioren": 80, "B-Juniorinnen": 80,
        "A-Junioren": 90, "A-Juniorinnen": 90,
    }

    belegung: dict[tuple[str, str], list[tuple[int, int, bool]]] = defaultdict(list)

    for plan in spielplaene:
        ak = plan.get("altersklasse", "")
        halbfeld = ak in _AK_HALBFELD
        dauer = _AK_DAUER.get(ak, 80)
        for st in plan.get("spieltage", []):
            for spiel in st.get("spiele", []):
                feld = (spiel.get("spielfeld") or "").strip().lower()
                datum = spiel.get("datum") or ""
                if not feld or not datum:
                    continue
                key = (feld, datum)
                zeit_str = spiel.get("anstosszeit", "") or ""
                parts = zeit_str.replace(":", ".").split(".")
                if len(parts) == 2 and parts[0].isdigit() and parts[1].isdigit():
                    start = int(parts[0]) * 60 + int(parts[1])
                else:
                    start = 720
                belegung[key].append((start, start + dauer, halbfeld))

    # Konflikte zählen pro (feld, datum)
    remaining: dict[tuple[str, str], int] = {}
    for key, entries in belegung.items():
        entries.sort()
        konflikte = 0
        for i in range(len(entries)):
            for j in range(i + 1, len(entries)):
                s_i, e_i, hf_i = entries[i]
                s_j, e_j, hf_j = entries[j]
                if s_j < e_i:
                    if hf_i and hf_j and s_i == s_j:
                        continue
                    konflikte += 1
        remaining[key] = konflikte

    # Verteile auf Spielpläne
    for plan in spielplaene:
        if not plan.get("score"):
            continue
        plan["score"]["platz_konflikte"] = 0
        ak = plan.get("altersklasse", "")
        for st in plan.get("spieltage", []):
            for spiel in st.get("spiele", []):
                feld = (spiel.get("spielfeld") or "").strip().lower()
                datum = spiel.get("datum") or ""
                if not feld or not datum:
                    continue
                key = (feld, datum)
                if remaining.get(key, 0) > 0:
                    plan["score"]["platz_konflikte"] += 1
                    remaining[key] -= 1


# ─── Spielplan Excel Export ────────────────────────────────────

@app.post("/api/spielplan/export")
async def api_spielplan_export(request: Request):
    """Generiert eine Excel-Datei aus den Spielplänen.

    Akzeptiert entweder:
      - {"spielplaene": [...]} direkt
      - {"from_log": "spielplan_20260420_205718.json"} zum Laden aus Log
      - {} (leer) -> lädt automatisch den neuesten Log
    """
    data = await request.json()
    spielplaene = data.get("spielplaene")

    if not spielplaene:
        # Versuche aus Log zu laden
        log_file = data.get("from_log")
        if log_file:
            fp = SPIELPLAN_DIR / log_file
        else:
            # Neuesten Log finden
            logs = sorted(SPIELPLAN_DIR.glob("spielplan_*.json"))
            fp = logs[-1] if logs else None

        if fp and fp.exists():
            with open(fp, "r", encoding="utf-8") as f:
                log_data = json.load(f)
            spielplaene = log_data.get("spielplaene", [])

    if not spielplaene:
        raise HTTPException(400, "Keine Spielpläne vorhanden")

    wb = Workbook()

    # ── Styles ──
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="C41230", end_color="C41230", fill_type="solid")
    spieltag_font = Font(bold=True, size=11, color="C41230")
    thin_border = Border(bottom=Side(style="thin", color="E5E7EB"))
    center = Alignment(horizontal="center")
    wrap = Alignment(wrap_text=True, vertical="center")

    # ── Sheet 1: Übersicht ──
    ws_ueb = wb.active
    ws_ueb.title = "Übersicht"
    ueb_headers = ["Altersklasse", "Topf", "Staffel", "Teams", "Spieltage",
                    "Spiele", "Konflikte", "Doppelrunde"]
    for ci, h in enumerate(ueb_headers, 1):
        cell = ws_ueb.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for ri, p in enumerate(spielplaene, 2):
        n_spiele = sum(len(st["spiele"]) for st in p["spieltage"])
        konflikte = p.get("score", {}).get("platz_konflikte", 0) if p.get("score") else 0
        ws_ueb.cell(row=ri, column=1, value=p["altersklasse"])
        ws_ueb.cell(row=ri, column=2, value=p["topf"])
        ws_ueb.cell(row=ri, column=3, value=p["staffel_name"])
        ws_ueb.cell(row=ri, column=4, value=p["n_teams"]).alignment = center
        ws_ueb.cell(row=ri, column=5, value=len(p["spieltage"])).alignment = center
        ws_ueb.cell(row=ri, column=6, value=n_spiele).alignment = center
        c = ws_ueb.cell(row=ri, column=7, value=konflikte)
        c.alignment = center
        if konflikte > 0:
            c.font = Font(bold=True, color="DC2626")
        ws_ueb.cell(row=ri, column=8, value="Ja" if p.get("doppelrunde") else "Nein").alignment = center

    for col in ws_ueb.columns:
        ws_ueb.column_dimensions[col[0].column_letter].width = 16

    # ── Sheets per Spielplan ──
    for p in spielplaene:
        sheet_name = f"{p['altersklasse']} {p['staffel_name']}"[:31]
        ws = wb.create_sheet(title=sheet_name)

        # SZ-Lookup: Mannschaftsname -> SZ-Nummer
        sz_lookup = {}
        for z in p.get("sz_zuordnungen", []):
            sz_lookup[z["mannschaft"]] = z["sz"]

        row = 1
        for st in p["spieltage"]:
            # Spieltag header
            datum_str = ""
            if st.get("datum"):
                try:
                    from datetime import date as _date
                    dt = _date.fromisoformat(st["datum"])
                    wochentage = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]
                    datum_str = f"{wochentage[dt.weekday()]}, {dt.strftime('%d.%m.%Y')}"
                except (ValueError, IndexError):
                    datum_str = st["datum"]

            label = f"Spieltag {st['nummer']}  —  {datum_str}"
            cell = ws.cell(row=row, column=1, value=label)
            cell.font = spieltag_font
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            row += 1

            # Table headers
            sp_headers = ["Zeit", "SZ", "Heim", "SZ", "Gast", "Spielort", "Datum"]
            for ci, h in enumerate(sp_headers, 1):
                cell = ws.cell(row=row, column=ci, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
            row += 1

            # Spiele
            for spiel in st["spiele"]:
                datum_spiel = ""
                if spiel.get("datum"):
                    try:
                        dt2 = _date.fromisoformat(spiel["datum"])
                        datum_spiel = dt2.strftime("%d.%m.%Y")
                    except ValueError:
                        datum_spiel = spiel["datum"]

                heim = spiel.get("heim", "")
                gast = spiel.get("gast", "")
                ws.cell(row=row, column=1, value=spiel.get("anstosszeit", "")).alignment = center
                ws.cell(row=row, column=2, value=sz_lookup.get(heim, "")).alignment = center
                ws.cell(row=row, column=3, value=heim)
                ws.cell(row=row, column=4, value=sz_lookup.get(gast, "")).alignment = center
                ws.cell(row=row, column=5, value=gast)
                ws.cell(row=row, column=6, value=spiel.get("spielfeld", ""))
                ws.cell(row=row, column=7, value=datum_spiel).alignment = center
                for ci in range(1, 8):
                    ws.cell(row=row, column=ci).border = thin_border
                row += 1

            # Spielfrei
            if st.get("spielfrei"):
                ws.cell(row=row, column=1, value=f"Spielfrei: {st['spielfrei']}")
                ws.cell(row=row, column=1).font = Font(italic=True, color="888888")
                row += 1

            row += 1  # empty row

        # Column widths
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 6
        ws.column_dimensions["C"].width = 28
        ws.column_dimensions["D"].width = 6
        ws.column_dimensions["E"].width = 28
        ws.column_dimensions["F"].width = 36
        ws.column_dimensions["G"].width = 14

    # Save to temp file
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name)
    tmp.close()

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="Spielplaene.xlsx",
    )


# ─── Schneller Wünsche-Parser (Regex, kein LLM) ───────────────

import re as _re

_WOCHENTAGE = {
    "montag": "Montag", "dienstag": "Dienstag", "mittwoch": "Mittwoch",
    "donnerstag": "Donnerstag", "freitag": "Freitag", "samstag": "Samstag",
    "sonntag": "Sonntag",
}

_SPERR_PATTERN = _re.compile(
    r"(?:nicht|kein|gesperrt|fällt aus|keine? (?:spiel|platz))"
    r".*?(\d{1,2})[./](\d{1,2})(?:[./](\d{2,4}))?",
    _re.IGNORECASE,
)

# Sperrtag auch wenn Datum VOR dem Keyword steht:
# "Am 15.11. kein Heimspiel", "15.11. gesperrt"
_SPERR_PATTERN2 = _re.compile(
    r"(?:am\s+)?(\d{1,2})[./](\d{1,2})(?:[./](\d{2,4}))?"
    r"\s*\.?\s*(?:kein|nicht|gesperrt|fällt|geht nicht|keine)",
    _re.IGNORECASE,
)

_DATUM_PATTERN = _re.compile(
    r"(\d{1,2})[./](\d{1,2})(?:[./](\d{2,4}))?",
)

# Zeit-Pattern: explizit "HH:MM" (mit Doppelpunkt) oder "HH.MM Uhr"
# NICHT "DD.MM" (Datum) matchen
_ZEIT_PATTERN = _re.compile(
    r"(\d{1,2}):(\d{2})\s*(?:uhr)?|(\d{1,2})\.(\d{2})\s*uhr",
    _re.IGNORECASE,
)

# Explizite Anstoß-/Uhrzeit-Keywords vor einer Zahl
_ZEIT_KEYWORD_PATTERN = _re.compile(
    r"(?:anstoß|anstoss|anstoßzeit|uhrzeit|spielbeginn|beginn|ab)\s*(?:um|ab|:)?\s*(\d{1,2})[:.:](\d{2})",
    _re.IGNORECASE,
)

_HEIM_KEYWORDS = _re.compile(
    r"heim|heimspiel|zu\s*hause",
    _re.IGNORECASE,
)

_AUSW_KEYWORDS = _re.compile(
    r"auswärts|ausw[aä]rts",
    _re.IGNORECASE,
)

_PLATZ_KEYWORDS = _re.compile(
    r"platz.*(?:teil|shar|gemeinsam|gleichzeitig)|gleich(?:en?)\s*platz",
    _re.IGNORECASE,
)


def _saison_start_year(saison: str) -> int | None:
    """Extrahiert das Startjahr aus z.B. '2025/26' -> 2025."""
    if not saison:
        return None
    try:
        return int(saison.split("/")[0])
    except (ValueError, IndexError):
        return None


def _resolve_year(month: int, saison: str) -> int:
    """Bestimmt das Jahr für einen Monat (Aug-Dez -> Startjahr, Jan-Jul -> Startjahr+1)."""
    start = _saison_start_year(saison)
    if start is None:
        from datetime import datetime
        now = datetime.now()
        start = now.year if now.month >= 8 else now.year - 1
    return start if month >= 8 else start + 1


def _parse_wuensche_fast(
    einteilung: dict,
    saison: str,
) -> dict[str, list[Wunsch]]:
    """Parst Wünsche-Freitext regelbasiert (schnell, ohne LLM).

    Erkennt: Sperrtage, Wochentag-Präferenzen, Uhrzeiten, Heim/Auswärts-Wünsche.
    """
    result: dict[str, list[Wunsch]] = {}

    for gruppe in einteilung.get("gruppen", []):
        for staffel in gruppe.get("staffeln", []):
            for team in staffel.get("teams", []):
                text = team.get("wuensche_text", "").strip()
                if not text:
                    continue

                mannschaft = team.get("mannschaft", "")
                wuensche: list[Wunsch] = []

                # Normalisiere: _x000D_ (Excel-Zeilenumbruch) -> Leerzeichen
                text_clean = text.replace("_x000D_", " ").replace("\r", " ").replace("\n", " ")

                # Sammle alle Datums-Positionen (um sie von Uhrzeit-Erkennung auszuschließen)
                datum_spans: list[tuple[int, int]] = []

                # ── Sperrtage (Pattern 1: "kein/nicht ... DD.MM.") ──
                for m in _SPERR_PATTERN.finditer(text_clean):
                    day, month = int(m.group(1)), int(m.group(2))
                    year_str = m.group(3)
                    if year_str:
                        year = int(year_str) if len(year_str) == 4 else 2000 + int(year_str)
                    else:
                        year = _resolve_year(month, saison)
                    try:
                        from datetime import date
                        datum = date(year, month, day)
                        wuensche.append(Wunsch(
                            kategorie=WunschKategorie.SPERRTAG,
                            prioritaet=WunschPrio.HART,
                            beschreibung=f"Gesperrt am {datum.strftime('%d.%m.%Y')}",
                            datum=datum.isoformat(),
                            original_text=text,
                        ))
                        datum_spans.append((m.start(1), m.end()))
                    except ValueError:
                        pass

                # ── Sperrtage (Pattern 2: "DD.MM. kein/nicht ...") ──
                for m in _SPERR_PATTERN2.finditer(text_clean):
                    day, month = int(m.group(1)), int(m.group(2))
                    year_str = m.group(3)
                    if year_str:
                        year = int(year_str) if len(year_str) == 4 else 2000 + int(year_str)
                    else:
                        year = _resolve_year(month, saison)
                    try:
                        from datetime import date
                        datum = date(year, month, day)
                        # Vermeid Duplikate
                        existing = {w.datum for w in wuensche if w.kategorie == WunschKategorie.SPERRTAG}
                        if datum.isoformat() not in existing:
                            wuensche.append(Wunsch(
                                kategorie=WunschKategorie.SPERRTAG,
                                prioritaet=WunschPrio.HART,
                                beschreibung=f"Gesperrt am {datum.strftime('%d.%m.%Y')}",
                                datum=datum.isoformat(),
                                original_text=text,
                            ))
                        datum_spans.append((m.start(1), m.end()))
                    except ValueError:
                        pass

                # ── Alle Datums-Positionen sammeln ──
                for m in _DATUM_PATTERN.finditer(text_clean):
                    day_val = int(m.group(1))
                    month_val = int(m.group(2))
                    if 1 <= day_val <= 31 and 1 <= month_val <= 12:
                        datum_spans.append((m.start(), m.end()))

                def _in_datum_span_approx(pos: int, spans: list[tuple[int, int]]) -> bool:
                    """Prüft ob position nahe (±5 Zeichen) eines Datums-Spans liegt."""
                    return any(s - 5 <= pos <= e + 5 for s, e in spans)

                # ── Wochentag-Präferenz ──
                text_lower = text_clean.lower()
                seen_wochentage = set()

                # Wochentage nur zählen, wenn sie NICHT direkt vor einem Datum stehen
                # (z.B. "Samstag, 01.11." ist ein Sperrtag, kein Wunschwochentag)
                for wt_key, wt_name in _WOCHENTAGE.items():
                    idx = text_lower.find(wt_key)
                    while idx >= 0:
                        # Prüfe ob nach dem Wochentag ein Datum folgt (innerhalb ~15 Zeichen)
                        after = text_lower[idx + len(wt_key):idx + len(wt_key) + 15]
                        in_datum = _in_datum_span_approx(idx, datum_spans)
                        has_date_after = _re.search(r'^\s*,?\s*\d{1,2}[./]\d{1,2}', after) is not None
                        if not in_datum and not has_date_after:
                            seen_wochentage.add(wt_name)
                            break
                        idx = text_lower.find(wt_key, idx + 1)

                # Wenn mehrere Wochentage: nur einen Wunsch mit dem übergeordneten Konzept
                # Aber im Solver werden alle als Alternativen behandelt
                # -> Speichere EINEN Wunsch pro Wochentag, aber in _update_wunsch_verletzungen
                #   zählen wir nur als Verletzung wenn der Tag auf KEINEM der gewünschten Tage liegt
                for wt_name in seen_wochentage:
                    wuensche.append(Wunsch(
                        kategorie=WunschKategorie.WOCHENTAG,
                        prioritaet=WunschPrio.WEICH,
                        beschreibung=f"Bevorzugt {wt_name}",
                        wochentag=wt_name,
                        original_text=text,
                    ))

                # ── Uhrzeit (nur wenn nicht innerhalb eines Datums) ──
                def _in_datum_span(pos: int) -> bool:
                    return any(s <= pos <= e for s, e in datum_spans)

                zeit_found = False
                # Erst explizite Keywords prüfen ("Anstoß 14:00")
                zeit_kw_match = _ZEIT_KEYWORD_PATTERN.search(text_clean)
                if zeit_kw_match:
                    h, m = int(zeit_kw_match.group(1)), int(zeit_kw_match.group(2))
                    if 8 <= h <= 21:
                        wuensche.append(Wunsch(
                            kategorie=WunschKategorie.ANSTOSSZEIT,
                            prioritaet=WunschPrio.WEICH,
                            beschreibung=f"Anstoß {h:02d}:{m:02d}",
                            uhrzeit=f"{h:02d}:{m:02d}",
                            original_text=text,
                        ))
                        zeit_found = True

                if not zeit_found:
                    zeit_match = _ZEIT_PATTERN.search(text_clean)
                    if zeit_match:
                        # Gruppen 1,2 für "HH:MM", Gruppen 3,4 für "HH.MM Uhr"
                        if zeit_match.group(1) is not None:
                            h, m = int(zeit_match.group(1)), int(zeit_match.group(2))
                        else:
                            h, m = int(zeit_match.group(3)), int(zeit_match.group(4))
                        if 8 <= h <= 21 and not _in_datum_span(zeit_match.start()):
                            wuensche.append(Wunsch(
                                kategorie=WunschKategorie.ANSTOSSZEIT,
                                prioritaet=WunschPrio.WEICH,
                                beschreibung=f"Anstoß {h:02d}:{m:02d}",
                                uhrzeit=f"{h:02d}:{m:02d}",
                                original_text=text,
                            ))

                # ── Heim/Auswärts-Beziehung ──
                if _HEIM_KEYWORDS.search(text_clean) and _AUSW_KEYWORDS.search(text_clean):
                    # "Wenn Heim, dann andere Auswärts" -> Platzsharing-Hinweis
                    wuensche.append(Wunsch(
                        kategorie=WunschKategorie.PLATZSHARING,
                        prioritaet=WunschPrio.WEICH,
                        beschreibung="Heim/Auswärts-Abstimmung gewünscht",
                        original_text=text,
                    ))

                # ── Fallback: wenn nichts erkannt, als Sonstiges speichern ──
                if not wuensche:
                    wuensche.append(Wunsch(
                        kategorie=WunschKategorie.SONSTIGES,
                        prioritaet=WunschPrio.WEICH,
                        beschreibung=text_clean[:200],
                        original_text=text,
                    ))

                result[mannschaft] = wuensche

    return result
