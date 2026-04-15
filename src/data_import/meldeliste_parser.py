"""Parser für DFBnet-Meldelisten (Excel-Export).

Unterstützt verschiedene Excel-Formate (24/25 und 25/26),
da die Spaltenreihenfolge je nach Export variieren kann.
Die Zuordnung erfolgt header-basiert.
"""

import json
from pathlib import Path

import openpyxl

from src.common.models import Mannschaft, Spielstaette

# Mapping: interner Feldname -> mögliche Header-Bezeichnungen in der Excel
COLUMN_ALIASES = {
    "verein_nr": ["V. Nr.", "V.Nr."],
    "vereinsname": ["Vereinsname"],
    "mannschaftsname": ["Mannschaftsname"],
    "altersklasse": ["MS-Art"],
    "spielklasse": ["Spielklasse"],
    "region": ["Region"],
    "bezirk_alt": ["Bezirk alt"],
    "ms_nr": ["MS-Nr.", "MS-Nr"],
    "staffel_grob": ["Staffel Grobeinteilung", "Grobeinteilung"],
    "staffel": ["Staffel"],
    "topf": ["Topf"],
    "spieltag": ["Spieltag"],
    "uhrzeit": ["Uhrzeit"],
    "wuensche": ["Wünsche"],
    "spielfeld_1": ["Spielfeld 1"],
    "spielfeld_2": ["Spielfeld 2"],
    "spielfeld_3": ["Spielfeld 3"],
    "spielfeld_4": ["Spielfeld 4"],
    "spielstaette_name": ["Spielstätte"],
    "strasse": ["Straße"],
    "plz": ["PLZ"],
    "ort": ["Ort"],
    "sz": ["SZ"],
}


def _find_header_row(ws) -> tuple[int, dict[str, int]]:
    """Findet die Header-Zeile und gibt das Column-Mapping zurück.

    Sucht nach der Zeile die 'Vereinsname' enthält.

    Returns:
        (header_row_index, {feldname: col_index})
    """
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        row_vals = [str(v).strip() if v else "" for v in row]
        if "Vereinsname" in row_vals:
            col_map = {}
            for feldname, aliases in COLUMN_ALIASES.items():
                for alias in aliases:
                    if alias in row_vals:
                        col_map[feldname] = row_vals.index(alias)
                        break
            return i, col_map

    raise ValueError("Header-Zeile mit 'Vereinsname' nicht gefunden")


def _parse_spielstaette_from_spielfeld(spielfeld_str: str) -> Spielstaette | None:
    """Parst Spielstätte aus dem Spielfeld-String (Format 24/25).

    Format: "Name, Straße, PLZ Ort-Stadtteil"
    """
    if not spielfeld_str or not spielfeld_str.strip():
        return None
    full = spielfeld_str.strip()
    parts = full.split(", ")
    if len(parts) >= 2:
        name = parts[0]
        adresse = ", ".join(parts[1:])
        return Spielstaette(name=name, adresse=adresse)
    return Spielstaette(name=full, adresse=full)


def _parse_spielstaette_from_columns(name: str, strasse: str, plz: str, ort: str) -> Spielstaette | None:
    """Parst Spielstätte aus separaten Spalten (Format 25/26)."""
    if not strasse and not plz:
        return None
    adresse_parts = []
    if strasse:
        adresse_parts.append(str(strasse).strip())
    if plz and ort:
        adresse_parts.append(f"{str(plz).strip()} {str(ort).strip()}")
    elif plz:
        adresse_parts.append(str(plz).strip())

    adresse = ", ".join(adresse_parts) if adresse_parts else ""
    return Spielstaette(name=str(name or "").strip(), adresse=adresse)


def _get(row: tuple, col_map: dict[str, int], feldname: str, default=None):
    """Liest einen Wert aus der Zeile anhand des Feldnamens."""
    idx = col_map.get(feldname)
    if idx is None or idx >= len(row):
        return default
    val = row[idx]
    return val if val is not None else default


def _topf_ableiten(ms_nr: int, topf_explicit: str | None) -> int:
    """Leitet den Topf ab.

    - Explizite Topf-Spalte hat Vorrang (25/26-Format)
    - Sonst: MS-Nr 1 = Topf 1, Rest = Topf 2
    """
    if topf_explicit:
        t = str(topf_explicit).strip()
        if t in ("1", "Topf 1"):
            return 1
        if t in ("2", "Topf 2"):
            return 2
    return 1 if ms_nr == 1 else 2


def parse_meldeliste(excel_path: str) -> list[Mannschaft]:
    """Parst eine DFBnet-Meldeliste und gibt eine Liste von Mannschaften zurück.

    Args:
        excel_path: Pfad zur Excel-Datei

    Returns:
        Liste von Mannschaft-Objekten
    """
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb[wb.sheetnames[0]]

    header_idx, col_map = _find_header_row(ws)

    mannschaften: list[Mannschaft] = []

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= header_idx:
            continue

        verein_nr = _get(row, col_map, "verein_nr")
        if not verein_nr:
            continue

        vereinsname = str(_get(row, col_map, "vereinsname", "")).strip()
        mannschaftsname = str(_get(row, col_map, "mannschaftsname", "")).strip()
        altersklasse = str(_get(row, col_map, "altersklasse", "")).strip()
        spielklasse = str(_get(row, col_map, "spielklasse", "")).strip()
        region = str(_get(row, col_map, "region", "")).strip()
        bezirk_alt = str(_get(row, col_map, "bezirk_alt", "")).strip()

        ms_nr_raw = _get(row, col_map, "ms_nr", 1)
        try:
            ms_nr = int(ms_nr_raw)
        except (ValueError, TypeError):
            ms_nr = 1

        topf_explicit = _get(row, col_map, "topf")
        topf = _topf_ableiten(ms_nr, topf_explicit)

        # Spielstätte: je nach Format
        if "strasse" in col_map and _get(row, col_map, "strasse"):
            spielstaette = _parse_spielstaette_from_columns(
                name=_get(row, col_map, "spielstaette_name", ""),
                strasse=_get(row, col_map, "strasse", ""),
                plz=_get(row, col_map, "plz", ""),
                ort=_get(row, col_map, "ort", ""),
            )
        else:
            spielfeld_1 = str(_get(row, col_map, "spielfeld_1", ""))
            spielstaette = _parse_spielstaette_from_spielfeld(spielfeld_1)

        wuensche_text = str(_get(row, col_map, "wuensche", "")).strip()

        mannschaften.append(
            Mannschaft(
                verein_nr=str(verein_nr).strip(),
                vereinsname=vereinsname,
                mannschaftsname=mannschaftsname,
                altersklasse=altersklasse,
                spielklasse=spielklasse,
                region=region,
                bezirk_alt=bezirk_alt,
                ms_nr=ms_nr,
                spielstaette=spielstaette,
                topf=topf,
                wuensche_text=wuensche_text,
            )
        )

    wb.close()
    return mannschaften


def verknuepfe_koordinaten(
    mannschaften: list[Mannschaft],
    cache_path: str | None = None,
) -> int:
    """Verknüpft Spielstätten mit Geokoordinaten aus dem Geocode-Cache.

    Args:
        mannschaften: Liste von Mannschaft-Objekten
        cache_path: Pfad zum geocode_cache.json (default: config/geocode_cache.json)

    Returns:
        Anzahl der erfolgreich verknüpften Spielstätten
    """
    if cache_path is None:
        cache_path = str(Path(__file__).resolve().parent.parent.parent / "config" / "geocode_cache.json")

    with open(cache_path, "r", encoding="utf-8") as f:
        cache = json.load(f)

    verknuepft = 0
    for m in mannschaften:
        if m.spielstaette and m.spielstaette.adresse:
            key = m.spielstaette.adresse.strip().lower()
            entry = cache.get(key)
            if entry and entry.get("lat") is not None:
                m.spielstaette.lat = entry["lat"]
                m.spielstaette.lon = entry["lon"]
                verknuepft += 1

    return verknuepft
